import json
from pathlib import Path
import subprocess
from openpyxl import load_workbook
from pypdf import PdfReader
from conftest import _resolve_cli


def test_installed_help_and_doctor_use_original_service(invoke, real):
    response = subprocess.run(_resolve_cli()+['--help'],capture_output=True,text=True,encoding='utf-8')
    assert response.returncode==0 and all(x in response.stdout for x in ['materials','expenses','drafts','reports'])
    result, body=invoke(real,'doctor')
    assert body['workbench_required'] is False and body['runtime']['workspace_id']==real['workspace_id']
    assert body['operations']['glm_calls']==body['operations']['jev_calls']==0
    assert '"token":' not in result.stdout


def test_real_project_and_material_results_reused_without_ai(invoke, real):
    _,body=invoke(real,'project','inspect')
    assert body['summary']==real['summary']
    args=['materials','prepare']
    for key in real['material_ids']:args+=['--id',key]
    _,body=invoke(real,*args,'--recognize-missing')
    assert body['reused'] and not body['needs_recognition']
    assert body['operations']['recognition_requests']==body['operations']['glm_calls']==body['operations']['jev_calls']==0
    assert {r['entry']['id']:r['entry']['fields'] for r in body['entries']}==real['material_fields']


def test_real_human_categories_remain_even_with_classify_missing(invoke, real):
    args=['expenses','classify']
    for key in real['expense_ids']:args+=['--id',key]
    _,body=invoke(real,*args,'--classify-missing')
    assert body['reused'] and all(r['protected'] for r in body['items'])
    assert {r['id']:r['category'] for r in body['items']}==real['categories']
    assert body['operations']['classification_requests']==body['operations']['jev_calls']==0


def test_real_existing_report_download_is_byte_identical_and_idempotent(invoke, real, tmp_path):
    _,listed=invoke(real,'reports','list');assert listed['reports']
    report=listed['reports'][0];out=tmp_path/'report'
    args=['reports','download','--report-id',report['id'],'--output-dir',out]
    _,dry=invoke(real,'--dry-run',*args);assert dry['dry_run'] and not out.exists()
    _,body=invoke(real,*args)
    assert body['historical_snapshot'] and body['operations']['report_render_requests']==0
    for row in body['artifacts']:
        assert Path(row['path']).read_bytes()==(Path(real['native_export_directory'])/row['name']).read_bytes()
        print('\nExisting original '+row['kind']+': '+row['path'])
    xlsx=next(r['path'] for r in body['artifacts'] if r['kind']=='xlsx')
    assert load_workbook(xlsx,data_only=True).sheetnames
    assert PdfReader(next(r['path'] for r in body['artifacts'] if r['kind']=='pdf')).pages
    before={p.name:(p.read_bytes(),p.stat().st_mtime_ns) for p in out.iterdir()}
    _,cached=invoke(real,*args)
    assert cached['operations']['downloads']==0 and before=={p.name:(p.read_bytes(),p.stat().st_mtime_ns) for p in out.iterdir()}
    _,offline=invoke(real,'results','show',out);assert offline['request']['report_id']==report['id']


def test_installed_repl_and_invalid_project_are_json(invoke, real):
    _,lines=invoke(real,input='project inspect\nquit\n')
    assert lines[0]['mode']=='repl' and lines[1]['summary']==real['summary']
    wrong={**real,'project':'nonexistent'}
    result,body=invoke(wrong,'materials','list',check=False)
    assert result.returncode==4 and body['error']['code']=='PROJECT_NOT_FOUND'


def test_native_cross_project_and_missing_recognition_dry_run(invoke, native):
    result,body=invoke(native,'materials','prepare','--id',native['foreign'],'--recognize-missing',check=False)
    assert result.returncode==4 and body['error']['code']=='PROJECT_BOUNDARY'
    _,saved=invoke(native,'materials','prepare','--id',native['voucher']);assert saved['reused']
    _,pending=invoke(native,'materials','prepare','--id',native['pending']);assert pending['needs_recognition']
    _,dry=invoke(native,'--dry-run','materials','prepare','--id',native['pending'],'--recognize-missing')
    assert dry['would_recognize']==[native['pending']] and dry['operations']['recognition_requests']==0


def test_native_receive_duplicate_does_not_recognize_or_move(invoke, native, tmp_path):
    source=next((tmp_path/'workspace/工作台/统一收件').glob('*.pdf'))
    # Choose an existing file whose native material belongs to this project.
    _,materials=invoke(native,'materials','list')
    by_digest={p.stem:p for p in source.parent.glob('*.pdf')}
    source=by_digest[next(r['digest'] for r in materials['materials'])]
    _,body=invoke(native,'materials','receive','--file',source)
    assert body['duplicate'] and body['operations']['recognition_requests']==0
    wrong={**native,'project':native['other_project']}
    result,body=invoke(wrong,'materials','receive','--file',source,check=False)
    assert result.returncode==4 and body['error']['code']=='PROJECT_BOUNDARY'


def test_native_suggestion_to_draft_replay_and_changed_request_conflict(invoke, native, tmp_path):
    _,response=invoke(native,'matches','suggest');suggestions=response['suggestions']['suggestions']
    candidate=next(s for s in suggestions if s['voucher_count']==s['payment_count']==1)
    request={'groups':[{'suggestion_id':candidate['id'],'snapshot_token':candidate['snapshot_token'],
                        'fields':{'project_type':'other','project_name':'CLI synthetic','amount':f"{candidate['amount_cents']/100:.2f}"}}]}
    file=tmp_path/'request.json';file.write_text(json.dumps(request),encoding='utf-8')
    args=['drafts','create','--request',file,'--request-id','cli-synthetic-001']
    _,dry=invoke(native,'--dry-run',*args);assert dry['dry_run'] and dry['operations']['draft_requests']==0
    _,created=invoke(native,*args)
    assert created['status']=='pending_review' and not created['reused']
    assert all(not r['verified'] and r['stage']=='draft' for r in created['result']['items'])
    _,replay=invoke(native,*args)
    assert replay['reused'] and [r['id'] for r in replay['result']['items']]==[r['id'] for r in created['result']['items']]
    request['groups'][0]['fields']['project_name']='Changed choice';file.write_text(json.dumps(request),encoding='utf-8')
    result,body=invoke(native,*args,check=False)
    assert result.returncode==4 and body['error']['code']=='HTTP_400'


def test_native_explicit_classification_uses_current_auto_route(invoke, native):
    _,before=invoke(native,'expenses','classify','--id',native['expense'])
    assert before['needs_classification'] and before['operations']['classification_requests']==0
    _,after=invoke(native,'expenses','classify','--id',native['expense'],'--classify-missing')
    assert after['operations']['classification_requests']==1 and after['operations']['jev_calls'] is None
    assert after['status']=='partial' and after['unresolved_classification']==[native['expense']]


def test_native_fresh_receive_then_explicit_local_prepare(invoke, native, tmp_path):
    source=next((tmp_path/'workspace/工作台/统一收件').glob('*.pdf'))
    newfile=tmp_path/'New synthetic.pdf';newfile.write_bytes(source.read_bytes()+b'\n% separate CLI receive fixture\n')
    _,planned=invoke(native,'--dry-run','materials','receive','--file',newfile)
    assert planned['dry_run'] and not planned.get('entry')
    _,received=invoke(native,'materials','receive','--file',newfile)
    assert not received['duplicate'] and received['entry']['status']=='received'
    assert received['operations']['recognition_requests']==0
    _,prepared=invoke(native,'materials','prepare','--id',received['entry']['id'],'--recognize-missing')
    assert prepared['operations']['recognition_requests']==1 and prepared['operations']['glm_calls'] is None
    assert prepared['entries'][0]['entry']['id']==received['entry']['id']
